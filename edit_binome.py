import os
import glob
import openpyxl

files = glob.glob("*.xlsx")
gi_files = [f for f in files if "Informatique" in f or "GL" in f]

if not gi_files:
    print("No GI file found!")
    exit(1)

gi_file = gi_files[0]
print("Modifying:", gi_file)

wb = openpyxl.load_workbook(gi_file)
ws = wb.active

cnes = []
for row in range(2, ws.max_row + 1):
    val = ws.cell(row=row, column=1).value
    if val:
        cnes.append((row, str(val).strip()))

if len(cnes) >= 4:
    # Binome 1
    row1, cne1 = cnes[0]
    row2, cne2 = cnes[1]
    ws.cell(row=row1, column=5).value = cne2
    ws.cell(row=row2, column=5).value = cne1
    print(f"Paired {cne1} and {cne2}")
    
    # Binome 2
    row3, cne3 = cnes[2]
    row4, cne4 = cnes[3]
    ws.cell(row=row3, column=5).value = cne4
    ws.cell(row=row4, column=5).value = cne3
    print(f"Paired {cne3} and {cne4}")

ws.cell(row=1, column=5).value = "CNE BINOME"

wb.save(gi_file)
print("Saved.")
